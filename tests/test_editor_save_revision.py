# SPDX-FileCopyrightText: 2026 Veyllo GmbH
# SPDX-License-Identifier: AGPL-3.0-or-later
# Additional permissions and terms under AGPL Section 7: see LICENSING.md
"""An editor save never overwrites silently.

The measured gaps: no save route checked a revision, so the agent's edit of an open file -
or another tab's save - was overwritten by the next save with nothing said; the Excel,
PowerPoint and Word saves wrote back only what the editor shows, so formulas, numbers, rows
past 500, pictures and unsupported Word content were lost from the original; a Markdown file
was sniffed for HTML tags and rewritten when it merely contained a table; the native Word
editor took an EMPTY model when its load failed and kept Save enabled. Each test names the
mutation it catches.
"""
import asyncio
import hashlib
import io
import os
import sys
import zipfile
from pathlib import Path

import pytest
from fastapi import HTTPException

from test_file_routes_account_jail import TENANT, _request, tree  # noqa: F401 - fixtures

ROOT = Path(__file__).resolve().parents[1]


def _rev(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _save(fn, body_cls, **body):
    return asyncio.run(fn(body_cls(**body), _request(TENANT, path="/api/file/save")))


def _conflict(fn, body_cls, **body):
    with pytest.raises(HTTPException) as refused:
        _save(fn, body_cls, **body)
    assert refused.value.status_code == 409
    return refused.value.detail


def test_the_file_route_names_the_revision(tree):
    from vaf.core import web_server as ws
    own = tree["own"]
    response = asyncio.run(ws.get_file(_request(TENANT), str(own)))
    assert response.headers["X-VAF-Revision"] == _rev(own)


def test_a_file_changed_since_it_was_opened_is_not_overwritten(tree):
    """MUTATION: drop the comparison in write_if_revision."""
    from vaf.core import web_server as ws
    own = tree["own"]
    opened = _rev(own)
    own.write_text("the agent's edit", encoding="utf-8")
    detail = _conflict(ws.save_file, ws.FileSaveRequest, path=str(own), content="mine",
                       base_revision=opened)
    assert detail["code"] == "conflict" and detail["current_revision"] == _rev(own)
    assert own.read_text(encoding="utf-8") == "the agent's edit"


def test_a_save_that_names_no_revision_does_not_overwrite_an_existing_file(tree):
    from vaf.core import web_server as ws
    detail = _conflict(ws.save_file, ws.FileSaveRequest, path=str(tree["own"]), content="mine")
    assert detail["code"] == "conflict"
    assert tree["own"].read_text(encoding="utf-8") == "x"


def test_the_right_revision_saves_and_answers_the_new_one(tree):
    from vaf.core import web_server as ws
    own = tree["own"]
    out = _save(ws.save_file, ws.FileSaveRequest, path=str(own), content="mine",
                base_revision=_rev(own))
    assert own.read_text(encoding="utf-8") == "mine" and out["revision"] == _rev(own)
    assert out["redirected"] is False


@pytest.mark.skipif(sys.platform == "win32", reason="POSIX file modes")
def test_a_save_keeps_the_files_mode(tree):
    """A document is not a secret. MUTATION: write it with the stores' 0600 default."""
    from vaf.core import web_server as ws
    own = tree["own"]
    os.chmod(own, 0o640)
    _save(ws.save_file, ws.FileSaveRequest, path=str(own), content="mine", base_revision=_rev(own))
    assert (own.stat().st_mode & 0o777) == 0o640


def test_markdown_is_written_verbatim_unless_the_editor_says_it_sends_html(tree):
    """MUTATION: sniff the content for tags again."""
    from vaf.core import web_server as ws
    own = tree["own"]
    raw = "| a | b |\n|---|---|\n\n<table><tr><td>kept as written</td></tr></table>\n"
    _save(ws.save_file, ws.FileSaveRequest, path=str(own), content=raw, base_revision=_rev(own))
    assert own.read_text(encoding="utf-8") == raw
    _save(ws.save_file, ws.FileSaveRequest, path=str(own), content="<h1>Title</h1><p>Body</p>",
          base_revision=_rev(own), format="html")
    assert own.read_text(encoding="utf-8").startswith("# Title")


def _workbook_with_formula(path: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws_ = wb.active
    ws_.title = "Sheet1"
    ws_["A1"], ws_["A2"], ws_["A3"] = "1", "2", "=A1+A2"
    wb.save(path)


def test_a_lossy_workbook_is_saved_to_one_copy_and_then_in_place(tree):
    """MUTATION: write a lossy office file in place (drop the loss check)."""
    pytest.importorskip("openpyxl")
    from vaf.core import web_server as ws
    original = tree["own"].with_name("budget.xlsx")
    _workbook_with_formula(original)
    before = original.read_bytes()

    loaded = asyncio.run(ws.get_file_as_html(_request(TENANT), str(original)))
    assert "formulas" in loaded.headers["X-VAF-Loss"]
    html = "<table><tr><td>1</td></tr><tr><td>2</td></tr><tr><td>3</td></tr></table>"
    first = _save(ws.save_file_as_xlsx, ws.FileSaveRequest, path=str(original), content=html,
                  base_revision=loaded.headers["X-VAF-Revision"])
    copy = original.with_name("budget (bearbeitet).xlsx")
    assert first["redirected"] is True and Path(first["path"]) == copy
    assert original.read_bytes() == before, "the original was overwritten"

    # The copy is the editor's own: nothing to lose, saved in place, no second copy.
    assert ws._office_loss_report(copy) == []
    second = _save(ws.save_file_as_xlsx, ws.FileSaveRequest, path=str(copy), content=html,
                   base_revision=first["revision"])
    assert second["redirected"] is False and Path(second["path"]) == copy
    assert sorted(p.name for p in original.parent.glob("budget*")) == [
        "budget (bearbeitet).xlsx", "budget.xlsx"]


def test_a_workbook_with_many_empty_sheets_is_judged_without_walking_them(tmp_path):
    """The loss report runs on the event loop; it walked 500 x 30 cells of every sheet, empty
    ones included, so 200 empty sheets held every request for seconds. MUTATION: drop the
    early answer for too many sheets, or walk past the used range again."""
    pytest.importorskip("openpyxl")
    import time

    import openpyxl

    from vaf.core import web_server as ws
    wb = openpyxl.Workbook()
    for i in range(1, 60):
        wb.create_sheet(f"Sheet{i + 1}")
    many = tmp_path / "many.xlsx"
    wb.save(many)
    started = time.monotonic()
    assert ws._office_loss_report(many) == ["too_large"]
    assert time.monotonic() - started < 1.0

    one = openpyxl.Workbook()
    one.active.title = "Sheet1"
    one.active["A1"] = "text"
    small = tmp_path / "small.xlsx"
    one.save(small)
    seen = []
    real_iter = openpyxl.worksheet.worksheet.Worksheet.iter_rows

    def counting(self, *a, **kw):
        for row in real_iter(self, *a, **kw):
            seen.extend(row)
            yield row

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(openpyxl.worksheet.worksheet.Worksheet, "iter_rows", counting)
        assert ws._office_loss_report(small) == []
    assert len(seen) == 1, f"walked {len(seen)} cells of a sheet that uses one"


def test_the_original_is_not_saved_over_an_existing_copy(tree):
    """A second session on the original: the copy exists, and it is opened, not replaced.
    MUTATION: write the copy whether or not it exists."""
    pytest.importorskip("openpyxl")
    from vaf.core import web_server as ws
    original = tree["own"].with_name("budget.xlsx")
    _workbook_with_formula(original)
    copy = original.with_name("budget (bearbeitet).xlsx")
    copy.write_bytes(b"earlier edits")
    loaded = asyncio.run(ws.get_file_as_html(_request(TENANT), str(original)))
    from urllib.parse import unquote
    assert Path(unquote(loaded.headers["X-VAF-Edit-Copy"])) == copy
    detail = _conflict(ws.save_file_as_xlsx, ws.FileSaveRequest, path=str(original),
                       content="<table><tr><td>x</td></tr></table>",
                       base_revision=loaded.headers["X-VAF-Revision"])
    assert detail["code"] == "copy_exists" and Path(detail["path"]) == copy
    assert copy.read_bytes() == b"earlier edits"


def test_a_word_document_with_content_the_editor_cannot_hold_goes_to_a_copy(tree):
    pytest.importorskip("docx")
    from docx import Document

    from vaf.core import web_server as ws
    original = tree["own"].with_name("letter.docx")
    plain = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Dear Alice")
    doc.save(plain)
    # The same document plus a comments part: the native model has no place for comments.
    with zipfile.ZipFile(io.BytesIO(plain.getvalue())) as src, zipfile.ZipFile(original, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("word/comments.xml", "<w:comments xmlns:w='x'/>")
    before = original.read_bytes()
    loaded = asyncio.run(ws.get_file_as_docx_model(_request(TENANT), str(original)))
    assert "comments" in loaded.headers["X-VAF-Loss"]
    out = _save(ws.save_file_as_docx_native, ws.FileSaveDocxNativeRequest, path=str(original),
                document={"title": "letter", "sections": []},
                base_revision=loaded.headers["X-VAF-Revision"])
    assert out["redirected"] is True and Path(out["path"]).name == "letter (bearbeitet).docx"
    assert original.read_bytes() == before


def test_the_dead_html_to_docx_route_is_gone():
    from vaf.core import web_server as ws
    assert not any(getattr(r, "path", "") == "/api/file/save-docx" for r in ws.app.routes)


# ── the editors ──────────────────────────────────────────────────────────────

def _src(rel: str) -> str:
    return (ROOT / rel).read_text(encoding="utf-8")


def test_a_failed_word_load_never_becomes_an_empty_document():
    """MUTATION: adopt createEmptyNativeDocx in the load's catch again, or enable Save."""
    src = _src("web/components/NativeDocxEditor.tsx")
    load = src[src.index("const loadModel = useCallback"):src.index("useEffect(() => {\n    if (documentModel || !filePath) return;")]
    assert "createEmptyNativeDocx" not in load and "setLoadFailed(true)" in load
    assert "disabled={isSaving || loadFailed || !documentModel}" in src
    # And it is SAID, with a retry, before the spinner that would otherwise spin forever.
    failed_view = src.index("if (loadFailed && !documentModel && !isLoading)")
    assert failed_view < src.index("if (isLoading || !documentModel)")
    assert "loadFailed onReload={() => void loadModel()}" in src[failed_view:failed_view + 1200]


def test_the_code_viewer_conflict_names_the_revision_on_disk_now():
    """A second change on disk under unsaved edits kept the first one's revision, so
    "Overwrite" was refused again. MUTATION: keep any existing conflict."""
    viewer = _src("web/components/CodeViewer.tsx")
    assert "prev.currentRevision === revision" in viewer


def test_the_code_viewer_forgets_the_previous_files_revision():
    """A file handed in directly, after a server file: the save named the server file's
    revision and its conflict banner stayed. MUTATION: drop the reset on open."""
    viewer = _src("web/components/CodeViewer.tsx")
    opened = viewer[viewer.index("// Initial load + live polling"):viewer.index("if (initialContent !== undefined) {")]
    assert "revisionRef.current = null;" in opened and "setConflict(null);" in opened


def test_an_agent_rewrite_keeps_a_dirty_draft():
    """MUTATION: reset the editor state on document_ready whatever it holds."""
    src = _src("web/app/page.tsx")
    ready = src[src.index("data.type === 'document_ready'"):src.index("data.type === 'editor_apply_edit'")]
    assert "prev.isOpen && prev.filePath === fp && prev.dirty" in ready
    assert "externalChange: true" in ready


def test_another_file_from_the_agent_waits_while_the_draft_is_unsaved():
    """The agent opening another file replaced a dirty draft without a word. It is offered in
    the editor's banner now, and the draft stays until the person says so.
    MUTATION: replace the state for another path whatever it holds."""
    src = _src("web/app/page.tsx")
    ready = src[src.index("data.type === 'document_ready'"):src.index("data.type === 'editor_apply_edit'")]
    assert "prev.isOpen && prev.dirty && prev.filePath && prev.filePath !== fp" in ready
    assert "pendingFile: { path: fp" in ready
    banner = _src("web/components/EditorFileBanner.tsx")
    assert "} else if (pendingFile) {" in banner and "t('openPending')" in banner
    for editor in ("web/components/DocumentEditor.tsx", "web/components/NativeDocxEditor.tsx"):
        assert "pendingFile={pendingFile}" in _src(editor), editor


def test_opening_the_same_file_again_loads_it_afresh():
    """The editor's key changes with the path or the nonce; the same path kept the old
    editor and its draft while the state said clean. MUTATION: leave the nonce as it was."""
    src = _src("web/app/page.tsx")
    opener = src[src.index("onOpenFile={(path) => setDocumentEditorState("):]
    opener = opener[:opener.index("}))}")]
    assert "loadNonce: prev.filePath === path ? (prev.loadNonce ?? 0) + 1 : 0" in opener


def test_every_editor_save_names_its_revision():
    for rel in ("web/components/DocumentEditor.tsx", "web/components/NativeDocxEditor.tsx",
                "web/components/CodeViewer.tsx"):
        src = _src(rel)
        assert "saveEditorFile(" in src and "base_revision:" in src, rel


def test_the_code_viewer_follows_a_save_into_the_copy():
    """After "keep mine as a copy" the viewer showed the original and called the edits saved:
    the next save went against the original's revision. MUTATION: drop the retarget."""
    viewer = _src("web/components/CodeViewer.tsx")
    redirect = viewer[viewer.index("if (outcome.redirected) {"):viewer.index("revisionRef.current = outcome.revision;")]
    assert "onRetarget(outcome.path)" in redirect and "setIsDirty(false)" not in redirect
    assert "onRetarget={(path) => setCodeViewerState(" in _src("web/app/page.tsx")


def test_the_document_editor_stays_unsaved_when_it_cannot_follow_the_copy():
    """A host without onRetarget: the edits are in the copy, so the draft must not read as
    saved against the original. MUTATION: mark the draft clean before the redirect check."""
    src = _src("web/components/DocumentEditor.tsx")
    save = src[src.index("setConflict(null);\n            const info: EditorFileInfo"):]
    stay = save[:save.index("savedContentRef.current = content;")]
    assert "outcome.redirected && outcome.path && !onRetarget" in stay and "return;" in stay
    assert "dirtyRef.current = false" not in stay


def test_the_word_editor_stays_unsaved_when_it_cannot_follow_the_copy():
    """The same rule in NativeDocxEditor. MUTATION: mark the model saved before the redirect
    check."""
    src = _src("web/components/NativeDocxEditor.tsx")
    save = src[src.index("const saveDocument = async"):]
    stay = save[:save.index("savedModelRef.current = documentModel;")]
    assert "outcome.redirected && outcome.path && !onRetarget" in stay and "return;" in stay
    assert "dirtyRef.current = false" not in stay


def test_each_revoke_button_names_what_it_revokes():
    """A screen reader heard "Revoke" on every row. MUTATION: label the button with the
    bare word again."""
    import json
    src = _src("web/components/settings/StandingGrantsSection.tsx")
    assert "aria-label={t('revokeItem', { name })}" in src
    assert "revoke({ tools: [name] }), name)" in src and "revoke({ dirs: [dir] }), dir)" in src
    for lang in ("en", "de", "ja", "ko", "th", "tr", "zh"):
        grants = json.loads(_src(f"web/messages/{lang}.json"))["grants"]
        assert "{name}" in grants["revokeItem"], lang


def _loop_stays_free(coro_factory):
    """Run the route while a ticker counts on the same loop: a route that blocks the loop for
    half a second lets it count nothing meanwhile."""
    async def go():
        ticks = 0

        async def ticker():
            nonlocal ticks
            while True:
                await asyncio.sleep(0.05)
                ticks += 1

        t = asyncio.create_task(ticker())
        try:
            await coro_factory()
        except Exception:
            pass
        t.cancel()
        return ticks

    return asyncio.run(go())


def test_opening_and_saving_an_office_file_leave_the_event_loop_free(tree, monkeypatch):
    """The loss report and the conversion parse the file on every open and save; on the event
    loop a large one held every other request. MUTATION: call them directly in the routes."""
    import time

    from vaf.core import web_server as ws
    monkeypatch.setattr(ws, "_office_loss_report", lambda target: time.sleep(0.5) or [])
    monkeypatch.setattr(ws, "_xlsx_to_html", lambda target: "<table></table>")
    book = tree["own"].with_name("slow.xlsx")
    book.write_bytes(b"not really a workbook")
    assert _loop_stays_free(lambda: ws.get_file_as_html(_request(TENANT), str(book))) >= 4

    monkeypatch.setattr(ws, "_render_xlsx", lambda html: time.sleep(0.5) or b"x")
    monkeypatch.setattr(ws, "_save_editor_file", lambda *a, **k: {"status": "ok"})
    request = ws.FileSaveRequest(path=str(book), content="<table></table>", base_revision=None)
    assert _loop_stays_free(lambda: ws.save_file_as_xlsx(
        request, _request(TENANT, path="/api/file/save"))) >= 4


def test_the_grants_list_shows_only_the_accounts_it_came_from():
    """On the render where an admin switches accounts, the previous list was still drawn with
    revoke buttons that post to the new account. MUTATION: render the loaded list whatever
    endpoint it came from."""
    src = _src("web/components/settings/StandingGrantsSection.tsx")
    assert "const data = loaded && loaded.endpoint === endpoint ? loaded.grants : null;" in src
    assert "setLoaded(next ? { endpoint: requested, grants: next } : null);" in src


def test_a_save_without_a_revision_reads_as_none():
    """MUTATION: turn a missing revision into an empty string."""
    src = _src("web/lib/editorFile.ts")
    assert "revision: payload.revision ? String(payload.revision) : null," in src
    assert "revision: string | null; redirected: boolean" in src


def test_the_turkish_revoke_label_needs_no_case_ending():
    import json
    assert json.loads(_src("web/messages/tr.json"))["grants"]["revokeItem"] == "Geri al: {name}"
